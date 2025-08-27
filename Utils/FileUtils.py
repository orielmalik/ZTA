import io
import vobject
from docx import Document
import openpyxl
import mailparser
from fastapi import UploadFile



async def process_file(file: UploadFile):
    filename = file.filename
    content_bytes = await file.read()

    if filename.endswith(".docx"):
        doc = Document(io.BytesIO(content_bytes))
        text = "\n".join([p.text for p in doc.paragraphs])
        file_type = "docx"

    elif filename.endswith(".xlsx"):
        wb = openpyxl.load_workbook(io.BytesIO(content_bytes))
        sheet = wb.active
        text = "\n".join(
            [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value]
        )
        file_type = "xlsx"

    elif filename.endswith(".vcf"):
        try:
            content_str = content_bytes.decode(errors="ignore")
            vcard = vobject.readOne(io.StringIO(content_str))
            text = getattr(vcard, "fn", None).value if hasattr(vcard, "fn") else ""
        except Exception:
            text = ""
        file_type = "vcf"

    elif filename.endswith(".eml"):
        mail = mailparser.parse_from_bytes(content_bytes)
        text = f"Subject: {mail.subject}\nFrom: {mail.from_}\nBody: {mail.body}"
        file_type = "eml"

    else:
        text = content_bytes.decode(errors="ignore")
        file_type = "text"

    return {"text": text, "type": file_type}
